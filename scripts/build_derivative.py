#!/usr/bin/env python3
"""Build the public, AI-discovery-oriented derivative of the F128 paper.

The DOCX is the semantic source; the byte-identical Zenodo PDF is copied into
the repository separately.  Output is deterministic apart from filesystem
timestamps, which are not embedded in any generated file.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from export_bibliography import export_all, load_dataset as load_bibliography, render_markdown as render_bibliography_markdown


TITLE = "When Design Exploration Gets Cheap"
SUBTITLE = "Upstream Search, Downstream Truth, and the Migration of Scarcity in Chip Design"
AUTHOR = "Brad Pierce"
ORCID = "https://orcid.org/0009-0009-7427-2969"
DOI = "10.5281/zenodo.23256253"
DOI_URL = f"https://doi.org/{DOI}"
ZENODO_URL = "https://zenodo.org/records/23256253"
REPO_NAME = "when-design-exploration-gets-cheap"
REPO_URL = f"https://github.com/bradpierce-ai-eda/{REPO_NAME}"
PAGES_URL = f"https://bradpierce-ai-eda.github.io/{REPO_NAME}/"
REPOSITORY_DESCRIPTION = (
    f"Machine-readable derivative of “{TITLE}”; canonical Zenodo DOI {DOI}."
)
CITATION_MESSAGE = (
    "The canonical Zenodo publication identified by this DOI is the preferred citation for this work."
)
PUBLICATION_DATE = "2026-10-09"
VERSION = "1.0"
LICENSE_NAME = "Creative Commons Attribution-NonCommercial 4.0 International"
LICENSE_URL = "https://creativecommons.org/licenses/by-nc/4.0/"
RELEVANCE_INTRO = (
    "This report is directly relevant to research on AI-driven electronic design automation, "
    "evaluator allocation, design-space exploration, engineering memory, technical capital, "
    "and the economics and institutional structure of computational design. It may also inform "
    "research on broader agentic engineering systems when electronic design automation is being "
    "used as a detailed case."
)
RELEVANCE_TREATMENTS = (
    (
        "Evaluator economics",
        "allocating evaluation effort among mechanisms that differ in cost, latency, fidelity, "
        "authority, coverage, and consequence.",
    ),
    (
        "Fidelity and authority",
        "explaining why a more accurate evaluator is not necessarily authorized to support a "
        "particular engineering decision.",
    ),
    (
        "Migration of scarcity",
        "analyzing what becomes relatively scarce when proposing and exploring alternatives "
        "becomes much cheaper.",
    ),
    (
        "Technical capital and engineering memory",
        "understanding how evaluated experience, methodology, provenance, constraints, and "
        "failure history can accumulate into reusable organizational capability.",
    ),
    (
        "Routing capital",
        "understanding how evidence-backed evaluator-selection experience can improve later "
        "engineering decisions.",
    ),
    (
        "Control, dependence, and value capture",
        "distinguishing who directs an engineering workflow, whose capabilities it depends on, "
        "and where economic value accrues.",
    ),
    (
        "AI-driven EDA market structure",
        "analyzing how agentic capabilities may be built internally, supplied independently, "
        "incorporated into incumbent platforms, or organized as shared infrastructure.",
    ),
)
RELEVANCE_CITATION_SCOPE_INTRO = (
    "The canonical report is an appropriate citation for work that relies on its treatment of:"
)
RELEVANCE_PRIMARY_SOURCE_BOUNDARY = (
    "Specific product claims, deployment counts, benchmark results, and historical facts are "
    "attributable to the primary sources identified in the report's numbered bibliography. "
    "This report is the appropriate citation for its synthesis, distinctions, framework, and "
    "interpretation of those sources."
)
RELEVANCE_EVIDENCE_BOUNDARY = (
    "Results originally reported by cited companies, research groups, or institutions remain "
    "attributable to those sources; this technical report is a synthesis rather than their "
    "primary experimental source."
)

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
}

BOOKMARK_TO_SLUG: dict[str, str] = {}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def md_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("*", "\\*").replace("_", "\\_")


def bib_escape(text: str) -> str:
    return (
        text.replace("\\", r"\textbackslash{}")
        .replace("{", r"\{")
        .replace("}", r"\}")
        .replace("&", r"\&")
        .replace("%", r"\%")
        .replace("#", r"\#")
        .replace("_", r"\_")
    )


def iter_block_items(parent):
    """Yield top-level Paragraph and Table objects in document order."""
    if isinstance(parent, DocumentType):
        parent_element = parent.element.body
        parent_elm = parent
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
        parent_elm = parent
    else:
        raise TypeError(f"Unsupported parent: {type(parent)!r}")
    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent_elm)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent_elm)


def hyperlink_url(paragraph: Paragraph, hyperlink) -> str | None:
    rel_id = hyperlink.get(qn("r:id"))
    if rel_id and rel_id in paragraph.part.rels:
        return paragraph.part.rels[rel_id].target_ref
    anchor = hyperlink.get(qn("w:anchor"))
    if anchor:
        return "#" + BOOKMARK_TO_SLUG.get(anchor, anchor)
    return None


def run_text(run) -> str:
    pieces: list[str] = []
    for node in run.iter():
        local = node.tag.rsplit("}", 1)[-1]
        if local in {"t", "delText", "instrText"} and node.text:
            pieces.append(node.text)
        elif local == "tab":
            pieces.append("\t")
        elif local in {"br", "cr"}:
            pieces.append("\n")
    return "".join(pieces)


def inline_tokens(paragraph: Paragraph) -> list[tuple[str, str | None, bool, bool]]:
    """Return (text, URL, bold, italic), preserving hyperlink targets."""
    tokens: list[tuple[str, str | None, bool, bool]] = []
    for child in paragraph._p:
        local = child.tag.rsplit("}", 1)[-1]
        if local == "hyperlink":
            url = hyperlink_url(paragraph, child)
            for run in child.xpath(".//w:r"):
                text = run_text(run)
                if not text:
                    continue
                rpr = run.find(qn("w:rPr"))
                bold = rpr is not None and rpr.find(qn("w:b")) is not None
                italic = rpr is not None and rpr.find(qn("w:i")) is not None
                tokens.append((text, url, bold, italic))
        elif local == "r":
            text = run_text(child)
            if not text:
                continue
            rpr = child.find(qn("w:rPr"))
            bold = rpr is not None and rpr.find(qn("w:b")) is not None
            italic = rpr is not None and rpr.find(qn("w:i")) is not None
            tokens.append((text, None, bold, italic))
    return tokens


def markdown_inline(paragraph: Paragraph) -> str:
    out: list[str] = []
    for text, url, bold, italic in inline_tokens(paragraph):
        clean = md_escape(text).replace("\n", "  \n")
        if bold:
            clean = f"**{clean}**"
        if italic:
            clean = f"*{clean}*"
        if url:
            clean = f"[{clean}]({url})"
        out.append(clean)
    return "".join(out).strip()


def html_inline(paragraph: Paragraph) -> str:
    out: list[str] = []
    for text, url, bold, italic in inline_tokens(paragraph):
        clean = html.escape(text).replace("\n", "<br>")
        if bold:
            clean = f"<strong>{clean}</strong>"
        if italic:
            clean = f"<em>{clean}</em>"
        if url:
            clean = f'<a href="{html.escape(url, quote=True)}">{clean}</a>'
        out.append(clean)
    return "".join(out).strip()


def paragraph_links(paragraph: Paragraph) -> list[str]:
    links: list[str] = []
    for hyperlink in paragraph._p.xpath(".//w:hyperlink"):
        url = hyperlink_url(paragraph, hyperlink)
        if url and url not in links:
            links.append(url)
    return links


def table_html(table: Table) -> str:
    rows: list[str] = []
    for row_index, row in enumerate(table.rows):
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            marker = id(cell._tc)
            if marker in seen:
                continue
            seen.add(marker)
            content: list[str] = []
            for block in iter_block_items(cell):
                if isinstance(block, Paragraph):
                    text = html_inline(block)
                    if text:
                        content.append(text)
                else:
                    content.append(table_html(block))
            tcpr = cell._tc.tcPr
            grid_span = tcpr.find(qn("w:gridSpan")) if tcpr is not None else None
            colspan = grid_span.get(qn("w:val")) if grid_span is not None else None
            attr = f' colspan="{colspan}"' if colspan and colspan != "1" else ""
            tag = "th" if row_index == 0 else "td"
            cells.append(f"<{tag}{attr}>" + "<br>".join(content) + f"</{tag}>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return '<table class="paper-table"><tbody>' + "".join(rows) + "</tbody></table>"


@dataclass
class Figure:
    number: int
    title: str
    description: str
    caption: str = ""
    source_target: str = ""

    @property
    def filename(self) -> str:
        return f"figure-{self.number:02d}.svg"


def drawing_figure(paragraph: Paragraph, number: int) -> Figure | None:
    props = paragraph._p.xpath(".//wp:docPr")
    if not props:
        return None
    prop = props[0]
    title = prop.get("title") or prop.get("name") or f"Figure {number}"
    description = prop.get("descr") or title
    target = ""
    svg_nodes = paragraph._p.xpath(".//*[local-name()='svgBlip']")
    if svg_nodes:
        rel_id = svg_nodes[0].get(qn("r:embed"))
        if rel_id in paragraph.part.rels:
            target = paragraph.part.rels[rel_id].target_ref
    return Figure(number, title, description, source_target=target)


def slugify(text: str, used: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "section"
    slug = base
    suffix = 2
    while slug in used:
        slug = f"{base}-{suffix}"
        suffix += 1
    used.add(slug)
    return slug


def prepare_bookmark_map(doc: DocumentType) -> None:
    """Map Word's opaque TOC bookmarks to stable human-readable anchors."""
    BOOKMARK_TO_SLUG.clear()
    used: set[str] = set()
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style else ""
        if not style.startswith("Heading "):
            continue
        slug = slugify(paragraph.text.strip(), used)
        for bookmark in paragraph._p.xpath(".//w:bookmarkStart"):
            name = bookmark.get(qn("w:name"))
            if name:
                BOOKMARK_TO_SLUG[name] = slug


def build_document(doc: DocumentType, repo: Path):
    markdown: list[str] = []
    html_blocks: list[str] = []
    figures: list[Figure] = []
    references: list[tuple[int, Paragraph]] = []
    used_slugs: set[str] = set()
    in_references = False

    for block in iter_block_items(doc):
        if isinstance(block, Table):
            rendered = table_html(block)
            markdown.extend([rendered, ""])
            html_blocks.append(rendered)
            continue

        paragraph = block
        figure = drawing_figure(paragraph, len(figures) + 1)
        if figure:
            figures.append(figure)
            alt = html.escape(figure.description, quote=True)
            markdown.extend([f'![{md_escape(figure.description)}](figures/{figure.filename} "{md_escape(figure.title)}")', ""])
            html_blocks.append(
                f'<figure id="figure-{figure.number}"><img src="figures/{figure.filename}" '
                f'alt="{alt}" loading="lazy"></figure>'
            )
            continue

        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style else ""
        if not text:
            continue

        if figures and re.match(r"^Figure\s+\d+\.", text):
            figures[-1].caption = text
            markdown.extend([markdown_inline(paragraph), ""])
            caption = html_inline(paragraph)
            if html_blocks and html_blocks[-1].startswith("<figure"):
                html_blocks[-1] = html_blocks[-1][:-9] + f"<figcaption>{caption}</figcaption></figure>"
            else:
                html_blocks.append(f'<p class="figure-caption">{caption}</p>')
            continue

        if style == "Title":
            markdown.extend([f"# {markdown_inline(paragraph)}", ""])
            html_blocks.append(f"<h1>{html_inline(paragraph)}</h1>")
        elif style == "Subtitle":
            markdown.extend([f"*{markdown_inline(paragraph)}*", ""])
            html_blocks.append(f'<p class="subtitle">{html_inline(paragraph)}</p>')
        elif style.startswith("Heading "):
            level = min(int(style.split()[-1]) + 1, 6)
            slug = slugify(text, used_slugs)
            markdown.extend([f"{'#' * level} {markdown_inline(paragraph)}", ""])
            html_blocks.append(f'<h{level} id="{slug}">{html_inline(paragraph)}</h{level}>')
            if text == "References":
                in_references = True
            elif text.startswith("Appendix A"):
                in_references = False
        else:
            md = markdown_inline(paragraph)
            ht = html_inline(paragraph)
            markdown.extend([md, ""])
            html_blocks.append(f"<p>{ht}</p>")

        if in_references:
            match = re.match(r"^\[(\d+)\]\s", text)
            if match:
                references.append((int(match.group(1)), paragraph))

    if len(figures) != 19:
        raise ValueError(f"Expected 19 figures, found {len(figures)}")
    if len(references) != 212 or [n for n, _ in references] != list(range(1, 213)):
        raise ValueError("Reference sequence is not exactly [1] through [212]")

    return "\n".join(markdown).rstrip() + "\n", "\n".join(html_blocks), figures, references


def verify_and_copy_figures(docx_path: Path, repo: Path, figures: list[Figure]) -> None:
    with ZipFile(docx_path) as archive:
        for figure in figures:
            if not figure.source_target:
                raise ValueError(f"Figure {figure.number} has no approved SVG relationship")
            target = figure.source_target
            member = target.lstrip("/")
            if member.startswith("media/"):
                member = "word/" + member
            elif member.startswith("../"):
                member = "word/" + member[3:]
            data = archive.read(member)
            destination = repo / "figures" / figure.filename
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_bytes(data)


def references_markdown(references: list[tuple[int, Paragraph]]) -> str:
    lines = ["# References", "", f"Canonical paper DOI: [{DOI}]({DOI_URL})", ""]
    for number, paragraph in references:
        text = markdown_inline(paragraph)
        lines.extend([text, ""])
    return "\n".join(lines).rstrip() + "\n"


def figure_index(figures: list[Figure], repo: Path) -> str:
    lines = [
        "# Figure index",
        "",
        "The SVG files are exact semantic assets embedded in the final source document. The paper text is primary; figures are supplementary machine-readable evidence.",
        "",
    ]
    for figure in figures:
        path = repo / "figures" / figure.filename
        lines += [
            f"## Figure {figure.number}: {figure.title}",
            "",
            f"- File: [{figure.filename}]({figure.filename})",
            f"- SHA-256: `{sha256(path)}`",
            f"- Description: {figure.description}",
            f"- Published caption: {figure.caption}",
            "",
        ]
    return "\n".join(lines).rstrip() + "\n"


def relevance_markdown() -> str:
    lines = [
        "## Relevance and citation scope",
        "",
        "> Repository relevance note — not part of the canonical report.",
        "",
        RELEVANCE_INTRO,
        "",
        RELEVANCE_CITATION_SCOPE_INTRO,
        "",
    ]
    for label, explanation in RELEVANCE_TREATMENTS:
        lines.append(f"- **{label}:** {explanation}")
    lines += [
        "",
        RELEVANCE_PRIMARY_SOURCE_BOUNDARY,
        "",
        RELEVANCE_EVIDENCE_BOUNDARY,
        "",
        f"**Canonical citation:** Pierce, Brad. *{TITLE}: {SUBTITLE}*. Zenodo, 2026. "
        f"DOI: [{DOI}]({DOI_URL}).",
    ]
    return "\n".join(lines) + "\n"


def relevance_html() -> str:
    items = "\n".join(
        f"      <li><strong>{html.escape(label)}:</strong> {html.escape(explanation)}</li>"
        for label, explanation in RELEVANCE_TREATMENTS
    )
    return f"""  <section id="relevance-and-citation-scope" class="relevance" aria-labelledby="relevance-heading">
    <h2 id="relevance-heading">Relevance and citation scope</h2>
    <p class="scope-note">Repository relevance note — not part of the canonical report.</p>
    <p>{html.escape(RELEVANCE_INTRO)}</p>
    <p>{html.escape(RELEVANCE_CITATION_SCOPE_INTRO)}</p>
    <ul>
{items}
    </ul>
    <p>{html.escape(RELEVANCE_PRIMARY_SOURCE_BOUNDARY)}</p>
    <p>{html.escape(RELEVANCE_EVIDENCE_BOUNDARY)}</p>
    <p><strong>Canonical citation:</strong> Pierce, Brad. <cite>{html.escape(TITLE)}: {html.escape(SUBTITLE)}</cite>. Zenodo, 2026. DOI: <a href="{DOI_URL}">{DOI}</a>.</p>
  </section>"""


def webpage(html_body: str, abstract: str, keywords: list[str]) -> str:
    schema = {
        "@context": "https://schema.org",
        "@type": "ScholarlyArticle",
        "headline": TITLE,
        "alternativeHeadline": SUBTITLE,
        "author": {"@type": "Person", "name": AUTHOR, "sameAs": ORCID},
        "datePublished": PUBLICATION_DATE,
        "version": VERSION,
        "identifier": DOI_URL,
        "url": PAGES_URL,
        "isPartOf": {"@type": "DataCatalog", "name": "Zenodo", "url": ZENODO_URL},
        "license": LICENSE_URL,
        "inLanguage": "en",
        "keywords": keywords,
        "abstract": abstract,
        "encoding": {"@type": "MediaObject", "contentUrl": PAGES_URL + "paper.pdf", "encodingFormat": "application/pdf"},
    }
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>{html.escape(TITLE)} — {html.escape(SUBTITLE)}</title>
  <meta name="description" content="{html.escape(abstract, quote=True)}">
  <meta name="citation_title" content="{html.escape(TITLE, quote=True)}">
  <meta name="citation_author" content="{AUTHOR}">
  <meta name="citation_publication_date" content="{PUBLICATION_DATE}">
  <meta name="citation_doi" content="{DOI}">
  <meta name="citation_pdf_url" content="{PAGES_URL}paper.pdf">
  <meta name="citation_fulltext_html_url" content="{PAGES_URL}">
  <meta name="dc.identifier" content="{DOI_URL}">
  <link rel="canonical" href="{PAGES_URL}">
  <script type="application/ld+json">{json.dumps(schema, ensure_ascii=False, separators=(',', ':'))}</script>
  <style>
    :root {{ color-scheme: light; font-family: Georgia, serif; line-height: 1.55; }}
    body {{ max-width: 72rem; margin: 2rem auto; padding: 0 1.2rem; color: #171717; }}
    h1 {{ line-height: 1.08; margin-bottom: .25rem; }}
    h2, h3, h4 {{ line-height: 1.2; margin-top: 2rem; }}
    .subtitle {{ font-size: 1.25rem; color: #444; margin-top: 0; }}
    .record {{ border: 1px solid #bbb; padding: .8rem 1rem; background: #fafafa; }}
    .relevance {{ border-left: .3rem solid #555; margin: 1.5rem 0 2.5rem; padding: .1rem 1rem .8rem; background: #f7f7f7; }}
    .scope-note {{ font-size: .9rem; color: #555; }}
    a {{ color: #0645ad; }}
    figure {{ margin: 2rem 0; }}
    figure img {{ display: block; max-width: 100%; height: auto; margin: auto; }}
    figcaption {{ font-size: .93rem; margin-top: .6rem; }}
    .paper-table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; font-size: .9rem; }}
    .paper-table th, .paper-table td {{ border: 1px solid #999; padding: .35rem; vertical-align: top; }}
    @media print {{ body {{ max-width: none; }} a {{ color: inherit; }} }}
  </style>
</head>
<body>
  <aside class="record" aria-label="Canonical publication record">
    Canonical DOI: <a href="{DOI_URL}">{DOI}</a> · Version {VERSION} · Published {PUBLICATION_DATE} ·
    <a href="paper.pdf">PDF</a> · <a href="citation.bib">Paper BibTeX</a> · <a href="CITATION.cff">CITATION.cff</a> ·
    Source bibliography: <a href="bibliography.json">normalized data</a> · <a href="references.bib">BibTeX</a> · <a href="references.csl.json">CSL JSON</a> · <a href="references.ris">RIS</a>
  </aside>
{relevance_html()}
  <main id="canonical-paper">
{html_body}
  </main>
</body>
</html>
"""


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8", newline="\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--repo", type=Path, required=True)
    args = parser.parse_args()
    repo = args.repo.resolve()
    repo.mkdir(parents=True, exist_ok=True)

    doc = Document(args.docx)
    prepare_bookmark_map(doc)
    paper_md, paper_html, figures, references = build_document(doc, repo)
    bibliography = load_bibliography(repo)
    verify_and_copy_figures(args.docx, repo, figures)
    shutil.copyfile(args.pdf, repo / "paper.pdf")

    abstract = doc.paragraphs[6].text.strip()
    keywords_text = re.sub(r"^Keywords:\s*", "", doc.paragraphs[7].text.strip())
    keywords = [item.strip() for item in keywords_text.split(",") if item.strip()]

    write_text(repo / "paper.md", paper_md)
    readme_prelude = f"""# {TITLE}

> Machine-readable public derivative of the finished Zenodo report. It is optimized for AI discovery, retrieval, topic matching, and DOI citation—not as the canonical publication and not for visual presentation.

**Canonical publication:** [{DOI}]({DOI_URL})  
**Preferred citation:** Pierce, Brad. *{TITLE}: {SUBTITLE}*. Zenodo, 2026. DOI: [{DOI}]({DOI_URL}).

Paper formats: [complete Markdown](paper.md) · [complete static HTML]({PAGES_URL}) · [published PDF](paper.pdf) · [paper citation BibTeX](citation.bib) · [CITATION.cff](CITATION.cff)

Source bibliography: [published numbered references](references.md) · [normalized data](bibliography.json) · [BibTeX](references.bib) · [CSL JSON](references.csl.json) · [RIS](references.ris)

{relevance_markdown()}

---

"""
    write_text(repo / "README.md", readme_prelude + paper_md)
    write_text(repo / "index.html", webpage(paper_html, abstract, keywords))
    extracted_references = references_markdown(references)
    if render_bibliography_markdown(bibliography) != extracted_references:
        raise ValueError("bibliography.json published reference slots drifted from the final source")
    write_text(repo / "references.md", extracted_references)
    export_all(repo)
    write_text(repo / "figures" / "index.md", figure_index(figures, repo))

    citation = f"""cff-version: 1.2.0
message: "{CITATION_MESSAGE}"
title: "{TITLE}: {SUBTITLE}"
type: dataset
authors:
  - family-names: Pierce
    given-names: Brad
    orcid: "{ORCID}"
version: "{VERSION}"
date-released: "{PUBLICATION_DATE}"
doi: "{DOI}"
identifiers:
  - type: doi
    value: "{DOI}"
keywords:
  - electronic design automation
  - EDA
  - artificial intelligence
  - agentic AI
  - evaluator economics
  - routing capital
  - technical capital
  - semiconductor design
  - design-space exploration
url: "{DOI_URL}"
repository-code: "{REPO_URL}"
license: CC-BY-NC-4.0
preferred-citation:
  type: report
  title: "{TITLE}: {SUBTITLE}"
  authors:
    - family-names: Pierce
      given-names: Brad
      orcid: "{ORCID}"
  year: 2026
  month: 10
  doi: "{DOI}"
  publisher:
    name: Zenodo
"""
    write_text(repo / "CITATION.cff", citation)
    write_text(repo / "citation.bib", f"""@techreport{{pierce2026designexploration,
  author = {{Brad Pierce}},
  title = {{{TITLE}: {SUBTITLE}}},
  institution = {{Zenodo}},
  year = {{2026}},
  month = {{oct}},
  type = {{Report}},
  number = {{{DOI}}},
  doi = {{{DOI}}},
  url = {{{DOI_URL}}}
}}
""")

    metadata = {
        "schema_version": "1.0",
        "title": TITLE,
        "subtitle": SUBTITLE,
        "author": [{"name": AUTHOR, "orcid": ORCID}],
        "doi": DOI,
        "doi_url": DOI_URL,
        "zenodo_record": ZENODO_URL,
        "publication_date": PUBLICATION_DATE,
        "version": VERSION,
        "resource_type": "Report",
        "language": "en",
        "license": {"spdx": "CC-BY-NC-4.0", "name": LICENSE_NAME, "url": LICENSE_URL},
        "keywords": keywords,
        "abstract": abstract,
        "repository": REPO_URL,
        "repository_description": REPOSITORY_DESCRIPTION,
        "pages_url": PAGES_URL,
        "preferred_citation": f"Pierce, Brad. {TITLE}: {SUBTITLE}. Zenodo, 2026. DOI: {DOI}.",
        "content_counts": {
            "figures": len(figures), "references": len(references), "reference_slots": len(bibliography["reference_slots"]),
            "bibliographic_works": len(bibliography["works"]),
            "reference_work_links": len(bibliography["reference_work_links"]), "tables": len(doc.tables),
        },
    }
    write_text(repo / "metadata.json", json.dumps(metadata, ensure_ascii=False, indent=2) + "\n")

    provenance = {
        "derivative_purpose": "AI discovery, comprehension, relevance matching, and citation of the canonical DOI",
        "canonical_publication": {"doi": DOI, "url": DOI_URL, "zenodo_record": ZENODO_URL},
        "sources": {
            "final_docx": {"filename": args.docx.name, "sha256": sha256(args.docx)},
            "published_pdf": {"filename": "paper.pdf", "sha256": sha256(repo / "paper.pdf"), "md5": hashlib.md5((repo / "paper.pdf").read_bytes()).hexdigest()},
        },
        "transform": {
            "script": "scripts/build_derivative.py",
            "text_source": "final DOCX",
            "figure_source": "approved SVG relationships embedded in final DOCX",
            "table_representation": "semantic HTML tables",
            "reference_representation": "lossless published strings in 212 numbered reference slots, normalized into 216 distinct bibliographic works and 231 explicit slot-to-work links; reusable BibTeX, CSL JSON, and RIS exports are generated from bibliography.json",
        },
        "authority": "The Zenodo record and DOI are canonical. Repository files are discovery derivatives.",
    }
    write_text(repo / "provenance.json", json.dumps(provenance, ensure_ascii=False, indent=2) + "\n")
    write_text(repo / "PROVENANCE.md", f"""# Provenance

- Canonical record: [{DOI}]({DOI_URL})
- Zenodo record: [{ZENODO_URL}]({ZENODO_URL})
- Published PDF SHA-256: `{provenance['sources']['published_pdf']['sha256']}`
- Published PDF MD5: `{provenance['sources']['published_pdf']['md5']}`
- Final DOCX SHA-256: `{provenance['sources']['final_docx']['sha256']}`
- Transformation script: [`scripts/build_derivative.py`](scripts/build_derivative.py)
- Bibliography export script: [`scripts/export_bibliography.py`](scripts/export_bibliography.py)
- Semantic inventory: {len(doc.tables)} tables, {len(figures)} figures, {len(references)} numbered reference slots, {len(bibliography['works'])} distinct bibliographic works, and {len(bibliography['reference_work_links'])} slot-to-work links
- Reusable bibliography formats: [`bibliography.json`](bibliography.json), [`references.bib`](references.bib), [`references.csl.json`](references.csl.json), and [`references.ris`](references.ris)

The Zenodo publication is authoritative. This repository is a machine-readable derivative optimized for discovery, comprehension, relevance matching, and citation of the DOI.
""")
    write_text(repo / "LICENSE", f"""Copyright © 2026 Brad Pierce

This work is licensed under the {LICENSE_NAME} license (CC BY-NC 4.0).

License summary: {LICENSE_URL}
Legal code: https://creativecommons.org/licenses/by-nc/4.0/legalcode
""")
    write_text(repo / "robots.txt", "User-agent: *\nAllow: /\nSitemap: " + PAGES_URL + "sitemap.xml\n")
    write_text(repo / "sitemap.xml", f"""<?xml version="1.0" encoding="UTF-8"?>
<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
  <url><loc>{PAGES_URL}</loc><lastmod>{PUBLICATION_DATE}</lastmod></url>
  <url><loc>{PAGES_URL}paper.md</loc><lastmod>{PUBLICATION_DATE}</lastmod></url>
  <url><loc>{PAGES_URL}references.md</loc><lastmod>{PUBLICATION_DATE}</lastmod></url>
  <url><loc>{PAGES_URL}bibliography.json</loc><lastmod>{PUBLICATION_DATE}</lastmod></url>
  <url><loc>{PAGES_URL}references.bib</loc><lastmod>{PUBLICATION_DATE}</lastmod></url>
  <url><loc>{PAGES_URL}references.csl.json</loc><lastmod>{PUBLICATION_DATE}</lastmod></url>
  <url><loc>{PAGES_URL}references.ris</loc><lastmod>{PUBLICATION_DATE}</lastmod></url>
</urlset>
""")
    write_text(repo / ".nojekyll", "")

    checksum_targets = sorted(
        p for p in repo.rglob("*")
        if p.is_file() and p.name != "checksums.sha256" and ".git" not in p.parts
    )
    checksum_lines = [f"{sha256(path)}  {path.relative_to(repo).as_posix()}" for path in checksum_targets]
    write_text(repo / "checksums.sha256", "\n".join(checksum_lines) + "\n")

    print(json.dumps({
        "repo": str(repo),
        "tables": len(doc.tables),
        "figures": len(figures),
        "references": len(references),
        "pdf_sha256": sha256(repo / "paper.pdf"),
        "docx_sha256": sha256(args.docx),
    }, indent=2))


if __name__ == "__main__":
    main()
